"""Excel attendance report builder, shared by the on-demand /report command,
the "👥 Кто на паре" live Excel snapshot and the automatic weekly report.

Long/flat layout on purpose: one row per (session, student) with the date,
weekday, pair, subject, type and teacher spelled out — a wide pivot with
abbreviated multi-line headers is compact but hard to read at a glance."""
from __future__ import annotations

import datetime as dt
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import config
from database import Database

HEADER_FILL = PatternFill(start_color="FF2F5597", end_color="FF2F5597", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
PRESENT_FILL = PatternFill(start_color="FFD9EAD3", end_color="FFD9EAD3", fill_type="solid")
ABSENT_FILL = PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid")
EXCUSED_FILL = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")
PENDING_FILL = PatternFill(start_color="FFEFEFEF", end_color="FFEFEFEF", fill_type="solid")
THIN_BORDER = Border(*(Side(style="thin", color="FFBFBFBF"),) * 4)

STATUS_DISPLAY = {
    "present": ("Присутствовал", PRESENT_FILL),
    "absent": ("Отсутствовал", ABSENT_FILL),
    "excused": ("Уважительная причина", EXCUSED_FILL),
}
PENDING_DISPLAY = ("Ожидается", PENDING_FILL)


def _style_header(ws: Worksheet, headers: list[str], widths: list[int]) -> None:
    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


async def build_excel_report(db: Database, date_from: dt.date, date_to: dt.date) -> str:
    """Builds an Excel attendance report for [date_from, date_to] and
    returns the path of the saved (temp) file — the caller deletes it once
    sent. Sessions that haven't closed yet show "Ожидается" rather than
    being marked absent, so this is safe to call mid-class."""
    sessions = await db.get_sessions_range(date_from.isoformat(), date_to.isoformat())
    students = await db.get_all_students()
    attendance_rows = await db.get_attendance_range(date_from.isoformat(), date_to.isoformat())
    grade_rows = await db.get_grades_range(date_from.isoformat(), date_to.isoformat())

    status_map: dict[tuple[str, int, int], str] = {}
    for row in attendance_rows:
        status_map[(row["date"], row["pair_number"], row["student_id"])] = row["status"]

    grade_map: dict[tuple[str, int, int], str] = {}
    for row in grade_rows:
        grade_map[(row["date"], row["pair_number"], row["student_id"])] = row["grade"]

    wb = Workbook()

    # ---------------------------------------------------------------
    # Sheet 1: one row per (pair, student) — the actual attendance log
    # ---------------------------------------------------------------
    ws = wb.active
    ws.title = "Посещаемость"
    headers = ["Дата", "День", "Пара", "Время", "Предмет", "Тип", "Преподаватель", "Студент", "Статус", "Оценка"]
    widths = [12, 14, 6, 13, 24, 14, 20, 30, 22, 10]
    _style_header(ws, headers, widths)

    r = 2
    for session in sessions:
        date_obj = dt.date.fromisoformat(session["date"])
        for student in students:
            status = status_map.get((session["date"], session["pair_number"], student.id))
            if status:
                label, fill = STATUS_DISPLAY[status]
            elif session["closed"]:
                label, fill = STATUS_DISPLAY["absent"]
            else:
                label, fill = PENDING_DISPLAY
            grade = grade_map.get((session["date"], session["pair_number"], student.id), "")

            values = [
                date_obj.strftime("%d.%m.%Y"),
                config.WEEKDAY_FULL_RU[session["weekday"]],
                session["pair_number"],
                f"{session['start_dt'][11:16]}–{session['end_dt'][11:16]}",
                session["subject"],
                session["class_type"],
                session["teacher"],
                student.full_name,
                label,
                grade,
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=value)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center" if col != 8 else "left", vertical="center")
                if col == 9:
                    cell.fill = fill
            r += 1

    # ---------------------------------------------------------------
    # Sheet 2: one row per student — totals for the period
    # ---------------------------------------------------------------
    ws2 = wb.create_sheet("Итоги")
    headers2 = ["ФИО", "Пар всего", "Присутствий", "Пропусков", "По уважит. причине", "% посещаемости", "Средний балл"]
    widths2 = [30, 11, 12, 11, 18, 15, 13]
    _style_header(ws2, headers2, widths2)

    total_sessions = len(sessions)
    r = 2
    for student in students:
        present = absent = excused = 0
        numeric_grades: list[float] = []
        for session in sessions:
            status = status_map.get((session["date"], session["pair_number"], student.id))
            if status == "present":
                present += 1
            elif status == "excused":
                excused += 1
            elif status == "absent" or (status is None and session["closed"]):
                absent += 1
            grade = grade_map.get((session["date"], session["pair_number"], student.id))
            if grade is not None:
                try:
                    numeric_grades.append(float(grade.replace(",", ".")))
                except ValueError:
                    pass
        attendance_pct = (present / total_sessions) if total_sessions else 0
        avg_grade = (sum(numeric_grades) / len(numeric_grades)) if numeric_grades else None

        values = [student.full_name, total_sessions, present, absent, excused, attendance_pct, avg_grade]
        for col, value in enumerate(values, start=1):
            cell = ws2.cell(row=r, column=col, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if col != 1 else "left", vertical="center")
        ws2.cell(row=r, column=6).number_format = "0%"
        ws2.cell(row=r, column=7).number_format = "0.00"
        r += 1

    tmp_dir = tempfile.gettempdir()
    file_name = f"attendance_{date_from.isoformat()}_{date_to.isoformat()}.xlsx"
    file_path = os.path.join(tmp_dir, file_name)
    wb.save(file_path)
    return file_path


async def build_session_excel_report(db: Database, session) -> str:
    """Single-sheet ФИО+статус list for one just-closed pair — the compact
    per-pair attachment sent automatically after each class, as opposed to
    the multi-column period report above."""
    rows = sorted(await db.get_attendance_for_session(session["id"]), key=lambda r: r["full_name"])
    grade_by_student = {g["student_id"]: g["grade"] for g in await db.get_grades_for_session(session["id"])}

    wb = Workbook()
    ws = wb.active
    ws.title = "Посещаемость"
    headers = ["ФИО", "Статус", "Оценка"]
    widths = [34, 24, 10]
    _style_header(ws, headers, widths)

    r = 2
    for row in rows:
        label, fill = STATUS_DISPLAY[row["status"]]
        grade = grade_by_student.get(row["student_id"], "")
        for col, value in enumerate([row["full_name"], label, grade], start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center")
            if col == 2:
                cell.fill = fill
        r += 1

    tmp_dir = tempfile.gettempdir()
    file_name = f"session_{session['id']}.xlsx"
    file_path = os.path.join(tmp_dir, file_name)
    wb.save(file_path)
    return file_path
